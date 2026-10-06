"""
Student Performance Ranking Dashboard
Libraries: pandas, matplotlib
"""

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

# ---------------------------------------------------
# 1. DATA (sample data banayi hai; apni CSV bhi use kar sakte ho)
#    CSV use karne ke liye: df = pd.read_csv("students.csv")
#    Columns chahiye: Name, Maths, Science, English, Hindi, Computer
# ---------------------------------------------------
np.random.seed(42)
names = [
    "Aarav", "Riya", "Kabir", "Ananya", "Vivaan", "Isha", "Aditya", "Meera",
    "Rohan", "Sneha", "Arjun", "Pooja", "Karan", "Diya", "Rahul", "Simran",
    "Yash", "Nisha", "Dev", "Tanvi", "Harsh", "Kavya", "Manav", "Sara",
    "Ishaan", "Priya", "Neel", "Aditi", "Veer", "Zoya",
]
subjects = ["Maths", "Science", "English", "Hindi", "Computer"]

df = pd.DataFrame(
    np.random.randint(35, 100, size=(len(names), len(subjects))),
    columns=subjects,
)
df.insert(0, "Name", names)

# ---------------------------------------------------
# 2. PROCESSING: total, percentage, rank, grade
# ---------------------------------------------------
df["Total"] = df[subjects].sum(axis=1)
df["Percentage"] = (df["Total"] / (len(subjects) * 100) * 100).round(2)
df["Rank"] = df["Total"].rank(ascending=False, method="min").astype(int)


def grade(p):
    if p >= 90:
        return "A+"
    elif p >= 80:
        return "A"
    elif p >= 70:
        return "B"
    elif p >= 60:
        return "C"
    elif p >= 50:
        return "D"
    return "F"


df["Grade"] = df["Percentage"].apply(grade)
df = df.sort_values("Rank").reset_index(drop=True)

print("\n===== RANKING TABLE =====")
print(df[["Rank", "Name", "Total", "Percentage", "Grade"]].to_string(index=False))

print("\n===== QUICK STATS =====")
print("Topper        :", df.loc[0, "Name"], f"({df.loc[0, 'Percentage']}%)")
print("Class Average :", round(df["Percentage"].mean(), 2), "%")
print("Pass %        :", round((df["Grade"] != "F").mean() * 100, 1), "%")
print("Best subject  :", df[subjects].mean().idxmax())
print("Weak subject  :", df[subjects].mean().idxmin())

# Excel report (openpyxl chahiye: pip install openpyxl)
with pd.ExcelWriter("student_ranking_report.xlsx", engine="openpyxl") as writer:
    df.to_excel(writer, sheet_name="Ranking", index=False)
    ws = writer.sheets["Ranking"]
    from openpyxl.styles import Font, PatternFill, Alignment
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E86AB")
        cell.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 13
    ws.freeze_panes = "A2"

# ---------------------------------------------------
# 3. DASHBOARD (2 x 3 charts)
# ---------------------------------------------------
fig, axes = plt.subplots(2, 3, figsize=(18, 10))
fig.suptitle("Student Performance Ranking Dashboard", fontsize=20, fontweight="bold")

# (a) Top 10 students
top10 = df.head(10).iloc[::-1]
axes[0, 0].barh(top10["Name"], top10["Percentage"], color="#2E86AB")
axes[0, 0].set_title("Top 10 Students")
axes[0, 0].set_xlabel("Percentage")
for i, v in enumerate(top10["Percentage"]):
    axes[0, 0].text(v + 0.3, i, f"{v}", va="center", fontsize=8)

# (b) Subject-wise average
avg = df[subjects].mean()
axes[0, 1].bar(avg.index, avg.values, color="#F18F01")
axes[0, 1].set_title("Subject-wise Average Marks")
axes[0, 1].set_ylabel("Marks")
for i, v in enumerate(avg.values):
    axes[0, 1].text(i, v + 0.5, f"{v:.1f}", ha="center", fontsize=9)

# (c) Grade distribution
order = ["A+", "A", "B", "C", "D", "F"]
gc = df["Grade"].value_counts().reindex(order).dropna()
axes[0, 2].pie(gc.values, labels=gc.index, autopct="%1.0f%%", startangle=90)
axes[0, 2].set_title("Grade Distribution")

# (d) Percentage histogram
axes[1, 0].hist(df["Percentage"], bins=8, color="#6A994E", edgecolor="white")
axes[1, 0].axvline(df["Percentage"].mean(), color="red", linestyle="--", label="Average")
axes[1, 0].set_title("Percentage Distribution")
axes[1, 0].set_xlabel("Percentage")
axes[1, 0].set_ylabel("No. of Students")
axes[1, 0].legend()

# (e) Subject-wise box plot
axes[1, 1].boxplot([df[s] for s in subjects], tick_labels=subjects)
axes[1, 1].set_title("Marks Spread by Subject")
axes[1, 1].set_ylabel("Marks")

# (f) Topper vs class average (line chart)
axes[1, 2].plot(subjects, df.loc[0, subjects].values, marker="o", label=f"Topper ({df.loc[0, 'Name']})")
axes[1, 2].plot(subjects, avg.values, marker="s", linestyle="--", label="Class Average")
axes[1, 2].set_title("Topper vs Class Average")
axes[1, 2].set_ylabel("Marks")
axes[1, 2].legend()

plt.tight_layout(rect=[0, 0, 1, 0.95])
plt.savefig("student_dashboard.png", dpi=150)
plt.show()
